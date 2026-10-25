"""Fetch Australian open data and write the aggregated snapshots in data/open/ (OD-6).

Run with `uv run vetagent-open-data` (all sources) or `uv run vetagent-open-data ticks vets`.
The app only ever reads the snapshots, so tests and the demo stay offline and repeatable.
Snapshots hold aggregates only: no row-level records, no owner details.

The summarise_* functions are pure: they turn a source's own format into a snapshot.
"""

import argparse
import csv
import io
import json
import shutil
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from collections.abc import Iterable, Iterator
from datetime import UTC, datetime
from pathlib import Path

from vetagent.integrations.ezyvet import load_config
from vetagent.integrations.open_data import OPEN_DIR

USER_AGENT = "vetagent-demo/0.1 (open-data refresh)"
CKAN = "https://data.gov.au/data/api/3/action/package_show?id="
REGISTRATIONS_DATASET = "tcc-animal-registrations"
APVMA_DATASET = "0de37904-43e0-4814-b21b-5b64fafefe6f"
ALA_SEARCH = "https://biocache-ws.ala.org.au/ws/occurrences/search"
OVERPASS = "https://overpass-api.de/api/interpreter"
OPEN_METEO = "https://api.open-meteo.com/v1/forecast"

MONTHS = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]  # fmt: skip

# PubCRIS host codes: "AD..." are dogs, "AF..." are cats.
HOST_PREFIX = {"AD": "dog", "AF": "cat"}

# Pest codes worth a plain name on a companion-animal label; others use PubCRIS's own words.
PEST_NAMES = {
    "UTIPA": "paralysis tick",
    "UTIP": "paralysis tick",
    "UTIPT": "paralysis tick",
    "VGP1": "fleas",
    "IWC": "heartworm",
    "IWC1": "heartworm",
    "PIWH1": "heartworm",
}


def plain_pest(code: str) -> str | None:
    """Flea codes all start "WFL" (flea, cat flea, flea larvae, ...)."""
    return "fleas" if code.startswith("WFL") else PEST_NAMES.get(code)


# Pure transforms ------------------------------------------------------------------------


def summarise_registrations(sheets: dict[str, Iterable[dict]]) -> dict:
    """OD-1: counts per suburb, desexed or not, and the most registered breeds."""
    suburbs: dict[str, Counter] = defaultdict(Counter)
    breeds = {"dog": Counter(), "cat": Counter()}
    for sheet, rows in sheets.items():
        species = "cat" if "cat" in sheet.lower() else "dog"
        for row in rows:
            suburb = (row.get("Suburb") or "").strip().title()
            if not suburb:
                continue
            suburbs[suburb][f"{species}s"] += 1
            if str(row.get("Desexed") or "").strip().lower() == "yes":
                suburbs[suburb][f"{species}s_desexed"] += 1
            breeds[species][(row.get("Breed") or "Unknown").strip()] += 1
    keys = ["dogs", "dogs_desexed", "cats", "cats_desexed"]
    return {
        "suburbs": [{"suburb": s, **{k: c[k] for k in keys}} for s, c in sorted(suburbs.items())],
        "top_breeds": {sp: c.most_common(8) for sp, c in breeds.items()},
    }


def summarise_apvma(
    products: Iterable[dict],
    uses: Iterable[dict],
    pests: Iterable[dict],
    wanted: dict[str, str],
) -> dict:
    """OD-2: the labels of the products the clinic sells, plus how many dog and cat products
    are registered against paralysis tick, fleas and heartworm.

    `wanted` maps the clinic's product id to an APVMA product number.
    """
    vet = {
        r["pcode"].strip(): r for r in products if r["typedesc"].strip().startswith(("VM", "VP"))
    }
    pest_desc = {r["pestcode"].strip(): r["pestdesc"].strip().lower() for r in pests}
    numbers = set(wanted.values())
    label_uses: dict[str, dict[str, set]] = defaultdict(lambda: defaultdict(set))
    registered: dict[tuple[str, str], set] = defaultdict(set)
    for use in uses:
        pcode, host, pest = use["pcode"].strip(), use["hostcode"].strip(), use["pestcode"].strip()
        species = HOST_PREFIX.get(host[:2])
        if pcode not in vet or species is None:
            continue
        plain = plain_pest(pest)
        if pcode in numbers:
            label_uses[pcode][species].add(plain or pest_desc.get(pest, pest.lower()))
        if plain:
            registered[(species, plain)].add(pcode)
    labels = []
    for product_id, number in sorted(wanted.items()):
        row = vet.get(number)
        if row is None:
            continue
        labels.append(
            {
                "product_id": product_id,
                "apvma_number": number,
                "name": row["fpname"].strip(),
                "registrant": row["sname"].strip(),
                "expires": _pubcris_date(row["expdate"]),
                "uses": {sp: sorted(p) for sp, p in sorted(label_uses[number].items())},
            }
        )
    return {
        "labels": labels,
        "registered": [
            {"species": sp, "pest": pest, "products": len(codes)}
            for (sp, pest), codes in sorted(registered.items())
        ],
    }


def _pubcris_date(text: str) -> str:
    """PubCRIS dates look like "30/06/2027 12:00:00 AM"."""
    return datetime.strptime(text.split()[0], "%d/%m/%Y").date().isoformat()


def summarise_ticks(state: dict, local: dict, radius_km: int) -> dict:
    """OD-3: paralysis tick records by month in the state, and how many near the clinic."""
    by_month = {str(i + 1): 0 for i in range(12)}
    for facet in state.get("facetResults", []):
        if facet["fieldName"] == "month":
            for item in facet["fieldResult"]:
                if item["label"] in MONTHS:
                    by_month[str(MONTHS.index(item["label"]) + 1)] = item["count"]
    return {
        "state_by_month": by_month,
        "local": {"radius_km": radius_km, "records": local.get("totalRecords", 0)},
    }


def summarise_vets(overpass: dict) -> dict:
    """OD-4: named vets with their location, phone, hours and emergency tag."""
    vets = []
    for element in overpass.get("elements", []):
        tags = element.get("tags", {})
        lat = element.get("lat", element.get("center", {}).get("lat"))
        lon = element.get("lon", element.get("center", {}).get("lon"))
        if not tags.get("name") or lat is None:
            continue
        phone = next((tags[k] for k in ("phone", "contact:phone", "phone:AU") if tags.get(k)), "")
        vets.append(
            {
                "name": tags["name"],
                "lat": round(lat, 5),
                "lon": round(lon, 5),
                "phone": phone,
                "emergency": tags.get("emergency") == "yes",
                "hours": tags.get("opening_hours", ""),
                "osm": f"{element['type']}/{element['id']}",
            }
        )
    return {"vets": sorted(vets, key=lambda v: v["name"])}


def summarise_weather(forecast: dict) -> dict:
    """OD-5: each forecast day's feels-like maximum."""
    daily = forecast.get("daily", {})
    days = zip(daily.get("time", []), daily.get("apparent_temperature_max", []), strict=False)
    return {"days": [{"day": d, "feels_like_max": t} for d, t in days if t is not None]}


# Fetching -------------------------------------------------------------------------------


def _get(url: str, data: bytes | None = None, timeout: int = 120) -> bytes:
    """Public APIs are often busy (Overpass at capacity, a reset connection): back off and retry."""
    request = urllib.request.Request(url, data=data, headers={"User-Agent": USER_AGENT})
    for wait in (5, 20, 60, None):
        try:
            with urllib.request.urlopen(request, timeout=timeout) as response:
                return response.read()
        except urllib.error.HTTPError as exc:
            if exc.code not in (429, 502, 503, 504) or wait is None:
                raise
        except urllib.error.URLError:
            if wait is None:
                raise
        time.sleep(wait)
    raise AssertionError("unreachable")


def _json(url: str, params: dict | list | None = None) -> dict:
    query = f"?{urllib.parse.urlencode(params, doseq=True)}" if params else ""
    return json.loads(_get(url + query))


def _resources(dataset: str) -> list[dict]:
    return _json(CKAN + dataset)["result"]["resources"]


def _download(url: str, folder: Path) -> Path:
    """Stream to disk: the PubCRIS use table is about 70 MB."""
    path = folder / Path(urllib.parse.urlparse(url).path).name
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=600) as response, path.open("wb") as out:
        shutil.copyfileobj(response, out)
    return path


def _csv_rows(path: Path) -> Iterator[dict]:
    with path.open(encoding="latin-1", newline="") as f:
        yield from csv.DictReader(f)


def fetch_registrations(config: dict, folder: Path) -> dict:
    from openpyxl import load_workbook

    xlsx = [r for r in _resources(REGISTRATIONS_DATASET) if r.get("format", "").upper() == "XLSX"]
    latest = max(xlsx, key=lambda r: r.get("created", ""))
    book = load_workbook(io.BytesIO(_get(latest["url"])), read_only=True)
    sheets = {}
    for sheet in book.worksheets:
        rows = sheet.iter_rows(values_only=True)
        head = [str(h or "").strip() for h in next(rows)]
        sheets[sheet.title] = [dict(zip(head, r, strict=False)) for r in rows]
    return _stamp(
        summarise_registrations(sheets),
        source=latest["name"],
        url=latest["url"],
        publisher="Townsville City Council",
        licence="CC BY 4.0",
    )


def fetch_apvma(config: dict, folder: Path) -> dict:
    urls = {r["name"]: r["url"] for r in _resources(APVMA_DATASET)}
    files = {name: _download(urls[name], folder) for name in ("product.csv", "produse.csv")}
    files["pest.csv"] = _download(urls["pest.csv"], folder)
    summary = summarise_apvma(
        _csv_rows(files["product.csv"]),
        _csv_rows(files["produse.csv"]),
        _csv_rows(files["pest.csv"]),
        config["open_data"]["apvma"],
    )
    return _stamp(
        summary,
        source="PubCRIS register of agricultural and veterinary chemical products",
        url=f"https://data.gov.au/data/dataset/{APVMA_DATASET}",
        publisher="Australian Pesticides and Veterinary Medicines Authority",
        licence="CC BY 3.0 AU",
    )


def fetch_ticks(config: dict, folder: Path) -> dict:
    area = config["open_data"]
    taxon = 'taxa:"Ixodes holocyclus"'
    common = [("q", taxon), ("pageSize", 0), ("facets", "month"), ("flimit", 20)]
    state = _json(ALA_SEARCH, [*common, ("fq", f'stateProvince:"{area["state"]}"')])
    radius = area["tick_radius_km"]
    near = [("lat", area["lat"]), ("lon", area["lon"]), ("radius", radius)]
    local = _json(ALA_SEARCH, [*common, *near])
    return _stamp(
        {"state": area["state"], **summarise_ticks(state, local, radius)},
        source="Ixodes holocyclus (Australian paralysis tick) occurrence records",
        url="https://biocache.ala.org.au/occurrences/search?q=taxa%3A%22Ixodes%20holocyclus%22",
        publisher="Atlas of Living Australia",
        licence="Per-record licences, mostly CC BY; see the ALA data providers",
    )


def fetch_vets(config: dict, folder: Path) -> dict:
    area = config["open_data"]
    metres = area["vet_radius_km"] * 1000
    query = (
        f'[out:json][timeout:60];nwr["amenity"="veterinary"]'
        f"(around:{metres},{area['lat']},{area['lon']});out center tags;"
    )
    body = urllib.parse.urlencode({"data": query}).encode()
    result = json.loads(_get(OVERPASS, data=body))
    return _stamp(
        {"radius_km": area["vet_radius_km"], **summarise_vets(result)},
        source="amenity=veterinary within the radius of the clinic",
        url="https://www.openstreetmap.org/copyright",
        publisher="© OpenStreetMap contributors",
        licence="ODbL 1.0",
    )


def fetch_weather(config: dict, folder: Path) -> dict:
    area = config["open_data"]
    params = {
        "latitude": area["lat"],
        "longitude": area["lon"],
        "daily": "apparent_temperature_max",
        "timezone": config["timezone"],
        "forecast_days": 7,
    }
    return _stamp(
        summarise_weather(_json(OPEN_METEO, params)),
        source="7-day forecast, apparent (feels-like) maximum temperature",
        url="https://open-meteo.com/",
        publisher="Open-Meteo.com",
        licence="CC BY 4.0; the free API is for non-commercial use",
    )


def _stamp(snapshot: dict, **source: str) -> dict:
    fetched = datetime.now(UTC).date().isoformat()
    return {"_source": {**source, "fetched": fetched}, **snapshot}


FETCHERS = {
    "registrations": fetch_registrations,
    "apvma": fetch_apvma,
    "ticks": fetch_ticks,
    "vets": fetch_vets,
    "weather": fetch_weather,
}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Refresh the open-data snapshots in data/open/.")
    parser.add_argument("sources", nargs="*", help=f"any of: {', '.join(FETCHERS)} (default: all)")
    parser.add_argument("--out", type=Path, default=OPEN_DIR)
    args = parser.parse_args(argv)
    unknown = set(args.sources) - set(FETCHERS)
    if unknown:
        parser.error(f"unknown source: {', '.join(sorted(unknown))}")
    config = load_config()
    args.out.mkdir(parents=True, exist_ok=True)
    failed = 0
    with tempfile.TemporaryDirectory() as tmp:
        for name in args.sources or FETCHERS:
            try:
                snapshot = FETCHERS[name](config, Path(tmp))
            except Exception as exc:  # one source failing shouldn't lose the others
                print(f"{name}: failed, snapshot left as it was ({exc})", file=sys.stderr)
                failed += 1
                continue
            path = args.out / f"{name}.json"
            path.write_text(json.dumps(snapshot, indent=1) + "\n", encoding="utf-8")
            print(f"{name}: wrote {path}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
